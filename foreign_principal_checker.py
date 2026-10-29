#!/usr/bin/env python3
"""
foreign_principal_checker.py

Given a list of domains and a list of privileged groups, report the Foreign
Security Principals (members from another trusted domain/forest) in those groups,
labelling each as a USER or a GROUP.

Finding: "Foreign Security Principals in Privileged Group Allows Cross-Domain
Compromise." A foreign user or group inside a highly privileged group is a
cross-domain compromise path.

Auth is automatic: your account lives in one domain and reaches the others over
their trusts. The tool works out how to authenticate to each domain's DC (its own
name, your home domain, or letting netexec discover it), and also supports a
stored netexec credential with --id.

Usage
-----
    ./foreign_principal_checker.py -d domains.txt --groups-file groups.txt -U user -P 'pass' \
        -o fsp_report.txt --csv fsp.csv

    ./foreign_principal_checker.py -d domains.txt --groups "Administrators,Domain Admins" --id 1 \
        -o fsp_report.txt --csv fsp.csv

domains.txt : one per line, 'domain' or 'domain,dc'.
groups.txt  : one group per line (names with spaces are fine, e.g. Domain Admins).

Limitation: reports DIRECT members of each group. SID -> name/type resolution
works only for domains that are in the list and reachable.
"""

import argparse
import csv
import re
import subprocess
import sys

try:
    import dns.resolver as _dnsresolver
    HAVE_DNS = True
except ImportError:
    HAVE_DNS = False

DEFAULT_PRIV_GROUPS = [
    "Administrators", "Domain Admins", "Enterprise Admins", "Schema Admins",
    "Backup Operators", "Account Operators", "Server Operators",
    "Print Operators", "Group Policy Creator Owners", "DnsAdmins",
]

# FSP member DN: CN=<SID>,CN=ForeignSecurityPrincipals,DC=...
FSP_DN = re.compile(
    r"CN=(S-1-[0-9]+(?:-[0-9]+)+),CN=ForeignSecurityPrincipals,"
    r"(?:DC=[^,\s]+,?)+", re.IGNORECASE)
SAM_RE = re.compile(r"sAMAccountName\s+(.+)")        # value may contain spaces

_UNREACHABLE_SIGS = (
    "error connecting", "connection refused", "connection error",
    "unable to connect", "can't contact", "cannot contact",
    "no route to host", "timed out", "timeout", "network is unreachable",
    "name or service not known", "could not resolve", "socket",
)


def detect_auth_status(text):
    if "[+]" in text:
        return "ok"
    fail = "\n".join(l for l in text.splitlines() if "[-]" in l).lower()
    if not fail:
        return "unknown"
    if any(s in fail for s in _UNREACHABLE_SIGS):
        return "unreachable"
    return "auth_failed"


def _nxc_creds_args(creds, realm):
    """netexec auth args. realm is the -d value, or None to omit -d. A stored
    credential id short-circuits everything (like `nxc ... --id N`)."""
    if creds.get("cred_id"):
        return ["-id", str(creds["cred_id"])]
    args = []
    if realm:
        args += ["-d", realm]
    if creds.get("kerberos"):
        args += ["-k"]
    if creds.get("nthash"):
        args += ["-u", creds["username"], "-H", creds["nthash"]]
    else:
        args += ["-u", creds["username"], "-p", creds["password"]]
    return args


# ---------------------------- DNS / DC resolution --------------------------

def _run_nslookup(args, timeout):
    try:
        return subprocess.run(["nslookup"] + args, capture_output=True,
                              text=True, timeout=timeout).stdout or ""
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return ""


def _resolve_ip(host, dns_server=None, timeout=5):
    import ipaddress
    try:
        ipaddress.ip_address(host)
        return host
    except ValueError:
        pass
    if HAVE_DNS:
        res = _dnsresolver.Resolver(configure=(dns_server is None))
        if dns_server:
            res.nameservers = [dns_server]
        res.lifetime = res.timeout = timeout
        for rt in ("A", "AAAA"):
            try:
                return str(list(res.resolve(host, rt))[0])
            except Exception:
                continue
    out = _run_nslookup([host] + ([dns_server] if dns_server else []), timeout)
    seen_name = False
    for line in out.splitlines():
        if line.strip().startswith("Name:"):
            seen_name = True
            continue
        if "Address:" in line and "#" not in line and seen_name:
            return line.split("Address:")[-1].strip()
    return None


def resolve_dc(domain, dns_server=None, timeout=5):
    if HAVE_DNS:
        res = _dnsresolver.Resolver(configure=(dns_server is None))
        if dns_server:
            res.nameservers = [dns_server]
        res.lifetime = res.timeout = timeout
        for q in (f"_ldap._tcp.dc._msdcs.{domain}", f"_ldap._tcp.{domain}"):
            try:
                ans = res.resolve(q, "SRV")
                best = sorted(ans, key=lambda r: (r.priority, -r.weight))[0]
                host = str(best.target).rstrip(".")
                return _resolve_ip(host, dns_server) or host
            except Exception:
                continue
    for srv in (f"_ldap._tcp.dc._msdcs.{domain}", f"_ldap._tcp.{domain}"):
        text = _run_nslookup(["-type=SRV", srv] + ([dns_server] if dns_server else []), timeout)
        for line in text.splitlines():
            if "service =" in line or "hostname" in line.lower():
                tok = line.strip().split()[-1].rstrip(".")
                if tok and not tok.replace(".", "").isdigit():
                    return _resolve_ip(tok, dns_server) or tok
    return None


# ---------------------------- auth realm per domain ------------------------

def _auth_probe(dc, realm, creds, nxc_bin, timeout):
    cmd = [nxc_bin, "smb", dc] + _nxc_creds_args(creds, realm)
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return detect_auth_status((p.stdout or "") + (p.stderr or ""))
    except subprocess.TimeoutExpired:
        return "unreachable"
    except FileNotFoundError:
        return "error"


def prepare_domains(domains, creds, home_domain, dns_server, nxc_bin, timeout):
    """Resolve a DC + a working auth realm per domain. Returns
    [(domain, dc, realm, status)]."""
    smoke_t = min(timeout, 20)
    using_id = bool(creds.get("cred_id"))

    resolved = []
    for dom, dc in domains:
        if dc is None:
            dc = resolve_dc(dom, dns_server)
            if dc:
                print(f"[*] {dom}: resolved DC {dc}", file=sys.stderr)
            else:
                print(f"[!] {dom}: no DC found via SRV; add a DC in the file or use "
                      f"--dns-server", file=sys.stderr)
                dc = dom
        resolved.append((dom, dc))

    if using_id:
        out = []
        print("[*] Pre-flight: testing each DC with the stored credential (--id) ...",
              file=sys.stderr)
        for dom, dc in resolved:
            st = _auth_probe(dc, None, creds, nxc_bin, smoke_t)
            print(f"      {dom:<28} {dc:<16} {st}", file=sys.stderr)
            out.append((dom, dc, None, st))
        return out

    print("[*] Pre-flight: finding how to authenticate to each domain ...",
          file=sys.stderr)
    prelim = {}
    for dom, dc in resolved:
        won, status = None, "unreachable"
        for realm in (dom, None):
            st = _auth_probe(dc, realm, creds, nxc_bin, smoke_t)
            if st == "ok":
                won, status = realm, "ok"
                break
            if st in ("auth_failed", "unknown"):
                status = "auth_failed"
        prelim[dom] = (dc, status, won)
        note = f"  (realm: {won if won else 'auto'})" if status == "ok" else ""
        print(f"      {dom:<28} {dc:<16} {status}{note}", file=sys.stderr)

    if not home_domain:
        for dom, _dc in resolved:
            if prelim[dom][1] == "ok":
                home_domain = dom
                break
        if home_domain:
            print(f"[*] Home domain: {home_domain} (reaching the others over their trust)",
                  file=sys.stderr)

    out = []
    for dom, _dc in resolved:
        dc, status, realm = prelim[dom]
        if status == "auth_failed" and home_domain and home_domain != dom:
            if _auth_probe(dc, home_domain, creds, nxc_bin, smoke_t) == "ok":
                status, realm = "ok", home_domain
                print(f"[*] {dom}: authenticated via home domain '{home_domain}' "
                      f"(trust path)", file=sys.stderr)
        out.append((dom, dc, realm, status))
    return out


# ---------------------------- queries --------------------------------------

DEBUG = False


def ldap_query(dc, realm, filt, attrs, creds, nxc_bin, timeout):
    cmd = [nxc_bin, "ldap", dc] + _nxc_creds_args(creds, realm) + \
          ["--query", filt, attrs]
    if DEBUG:
        print(f"\n[debug] CMD: {' '.join(cmd)}", file=sys.stderr)
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        out = (p.stdout or "") + (p.stderr or "")
    except subprocess.TimeoutExpired:
        return "__ERROR__ timed out"
    except FileNotFoundError:
        return f"__ERROR__ '{nxc_bin}' not found"
    if DEBUG:
        print("[debug] RAW OUTPUT:\n" + (out.rstrip() or "(empty)"), file=sys.stderr)
        print(f"[debug] FSP DNs matched: {FSP_DN.findall(out) or 'none'}\n", file=sys.stderr)
    return out


def find_fsps_in_group(dc, realm, group, creds, nxc_bin, timeout):
    """Foreign SIDs that are direct members of a group.
    Returns (list_of_SIDs, error_or_None)."""
    out = ldap_query(
        dc, realm,
        f"(&(objectClass=group)(|(sAMAccountName={group})(cn={group})))",
        "member", creds, nxc_bin, timeout)
    if out.startswith("__ERROR__"):
        return [], out.replace("__ERROR__", "").strip()
    if detect_auth_status(out) in ("auth_failed", "unreachable") and not FSP_DN.search(out):
        return [], "cannot auth / unreachable"
    sids = []
    for m in FSP_DN.finditer(out):
        if m.group(1) not in sids:
            sids.append(m.group(1))
    return sids, None


def _attr_values(out, attr):
    vals = []
    for line in out.splitlines():
        mm = re.search(rf"(?:^|\s){re.escape(attr)}\s+(.+)$", line)
        if mm:
            vals.append(mm.group(1).strip())
    return vals


def resolve_principal(sid, scanned, creds, nxc_bin, timeout):
    """Resolve a foreign SID to 'DOMAIN\\name' and whether it is a USER or GROUP,
    by querying each in-scope domain. Returns (name_or_None, kind) where kind is
    'user' / 'group' / 'unknown' (unknown = resolved the name but not the type)."""
    for dom, dc, realm, status in scanned:
        if status != "ok":
            continue
        out = ldap_query(dc, realm, f"(objectSid={sid})",
                         "sAMAccountName objectClass objectCategory",
                         creds, nxc_bin, timeout)
        m = SAM_RE.search(out)
        if not m:
            continue
        oc = " ".join(_attr_values(out, "objectClass")).lower()
        cat = " ".join(_attr_values(out, "objectCategory")).lower()
        if "group" in oc or "cn=group" in cat:
            kind = "group"
        elif "user" in oc or "person" in cat:
            kind = "user"
        else:
            kind = "unknown"
        return f"{dom}\\{m.group(1).strip()}", kind
    return None, "unknown"


# ---------------------------- IO -------------------------------------------

def _kind_word(k):
    return {"group": "Foreign Group", "user": "Foreign User"}.get(
        k, "Foreign Principal (type undetermined)")


def _meaning(kind, grp, dom):
    if kind == "group":
        return f"Every member of this external group effectively has '{grp}' rights in {dom}."
    if kind == "user":
        return f"This external account has '{grp}' rights in {dom}."
    return f"This external principal is a member of '{grp}' in {dom}."


_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe(v):
    """Make a value safe for an Excel cell (strip control chars openpyxl rejects)."""
    return _ILLEGAL_XLSX.sub("", v) if isinstance(v, str) else v


def write_xlsx(path, scanned, groups, findings, resolved):
    """Colored Excel report: findings grouped by domain/group, principal type
    color-coded (red = foreign group / bigger blast radius, orange = foreign user)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    import datetime

    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    TYPE_FILL = {"group": ("FFC7CE", "9C0006"), "user": ("FFCC99", "974706")}
    GRAY = ("D9D9D9", "808080")
    arial = "Arial"

    total = 0
    for dom, dc, realm, status in scanned:
        for grp in groups:
            f = findings.get((dom, grp), {})
            if not f.get("error"):
                total += len(f.get("sids") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "FSP"

    ws["A1"] = "FOREIGN SECURITY PRINCIPALS IN PRIVILEGED GROUPS"
    ws["A1"].font = Font(name=arial, bold=True, size=13)
    ws["A2"] = ("External accounts/groups in privileged groups = cross-domain compromise "
                "paths. A foreign GROUP means everyone inside it inherits the access.")
    ws["A2"].font = Font(name=arial, italic=True, size=10)
    ws["A3"] = (f"Generated {now}   |   "
                f"Domains: {', '.join(d for d, _dc, _r, _s in scanned)}   |   "
                f"Groups: {', '.join(groups)}   |   {total} foreign principal(s) found")
    ws["A3"].font = Font(name=arial, size=10)

    header_row = 5
    headers = ["Domain", "Privileged Group", "Principal Type", "Foreign Principal",
               "What It Means", "Foreign SID"]
    hdr_fill = PatternFill("solid", fgColor="305496")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = Font(name=arial, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border

    NOTE_FILL = "F2F2F2"      # light gray for informational rows
    r = header_row + 1

    def write_row(values, type_fill=None, note=False):
        nonlocal r
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=_safe(v))
            cell.font = Font(name=arial, italic=note, color="808080" if note else "000000")
            cell.border = border
            if note:
                cell.fill = PatternFill("solid", fgColor=NOTE_FILL)
        if type_fill:
            fill, fontcolor = type_fill
            tcell = ws.cell(row=r, column=3)
            tcell.fill = PatternFill("solid", fgColor=fill)
            tcell.font = Font(name=arial, bold=True, color=fontcolor)
        r += 1

    for dom, dc, realm, status in scanned:
        # Domain we couldn't authenticate to / reach -- one clean note, like stdout.
        if status not in ("ok", "unknown"):
            write_row([dom, "", "Not checked", "",
                       f"Could not check this domain ({status}).", ""], note=True)
            continue

        shown = False
        for grp in groups:
            f = findings.get((dom, grp), {})
            if f.get("error"):
                # readable wording, matching the text report (not a stark QUERY ERROR)
                write_row([dom, grp, "Not checked", "",
                           f"Could not read this group ({f['error']}).", ""], note=True)
                shown = True
                continue
            for s in (f.get("sids") or []):
                shown = True
                name, kind = resolved.get(s, (None, "unknown"))
                write_row([dom, grp, _kind_word(kind), name or "(unresolved)",
                           _meaning(kind, grp, dom), s],
                          type_fill=TYPE_FILL.get(kind, GRAY))
        if not shown:
            # domain was checked cleanly, nothing foreign -- say so, like stdout
            write_row([dom, "", "Clean", "",
                       "No foreign principals found in the checked groups.", ""], note=True)

    for col, width in zip("ABCDEF", [26, 24, 20, 34, 60, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:F{max(r - 1, header_row)}"
    wb.save(path)


def read_lines(path):
    with open(path) as f:
        return [ln.strip() for ln in f if ln.strip() and not ln.startswith("#")]


def parse_domains(lines):
    out = []
    for ln in lines:
        if "," in ln:
            dom, dc = (p.strip() for p in ln.split(",", 1))
        else:
            dom, dc = ln, None
        out.append((dom, dc))
    return out


def main():
    ap = argparse.ArgumentParser(
        description="Report user and group Foreign Security Principals in the given "
                    "privileged groups, across a list of domains.",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-d", "--domains", required=True,
                    help="file of domains, one per line ('domain' or 'domain,dc')")
    ap.add_argument("--groups-file", help="file of groups to check, ONE PER LINE")
    ap.add_argument("--groups", help="comma-separated groups to check")
    ap.add_argument("-U", "--username", default="", help="account username")
    ap.add_argument("-P", "--password", default="", help="account password")
    ap.add_argument("-o", "--out", default="fsp_report.txt", help="text report path")
    ap.add_argument("--csv", help="also write a CSV here")
    ap.add_argument("--xlsx", help="also write a colored Excel report here "
                    "(red = foreign group, orange = foreign user)")

    g = ap.add_argument_group("auth (optional)")
    g.add_argument("--id", dest="cred_id", default="",
                   help="use a stored netexec credential id instead of -U/-P")
    g.add_argument("-H", "--hashes", dest="nthash", default="", help="NT hash for PtH")
    g.add_argument("-k", "--kerberos", action="store_true", help="use Kerberos")
    g.add_argument("--auth-domain", help="your home domain (auto-detected by default)")

    adv = ap.add_argument_group("advanced (optional)")
    adv.add_argument("--dns-server", metavar="IP", help="DNS server for SRV lookups")
    adv.add_argument("--no-resolve", action="store_true",
                     help="don't resolve SIDs to names/types (faster)")
    adv.add_argument("--nxc-bin", default="netexec", help="netexec binary name")
    adv.add_argument("--timeout", type=int, default=60, help="per-query timeout (s)")
    adv.add_argument("--threads", type=int, default=8,
                     help="parallel group/resolve queries (default 8)")
    adv.add_argument("--debug", action="store_true",
                     help="print each LDAP query command, its raw netexec output, and the "
                          "FSP DNs matched -- use to verify parsing against your version")
    args = ap.parse_args()

    global DEBUG
    DEBUG = args.debug

    domains = parse_domains(read_lines(args.domains))
    creds = {"username": args.username, "password": args.password,
             "nthash": args.nthash, "kerberos": args.kerberos, "cred_id": args.cred_id}

    # group list: file and/or comma; falls back to the built-in tier-0 list
    groups = []
    if args.groups_file:
        groups += read_lines(args.groups_file)
    if args.groups:
        groups += [x.strip() for x in args.groups.split(",") if x.strip()]
    if not groups:
        groups = list(DEFAULT_PRIV_GROUPS)
    seen, gl = set(), []
    for gg in groups:
        if gg.lower() not in seen:
            seen.add(gg.lower())
            gl.append(gg)
    groups = gl

    scanned = prepare_domains(domains, creds, args.auth_domain, args.dns_server,
                              args.nxc_bin, args.timeout)

    from concurrent.futures import ThreadPoolExecutor, as_completed
    findings, resolved = {}, {}

    # Build one task per (domain, group) across all checkable domains, and run
    # them in parallel -- same speedup model as pwdlast-check.
    tasks = []
    for dom, dc, realm, status in scanned:
        if status not in ("ok", "unknown"):
            print(f"[*] {dom}: {status} -- skipping", file=sys.stderr)
            continue
        for grp in groups:
            tasks.append((dom, dc, realm, grp))

    def run_group(t):
        dom, dc, realm, grp = t
        sids, err = find_fsps_in_group(dc, realm, grp, creds, args.nxc_bin, args.timeout)
        return (dom, grp), {"sids": sids, "error": err}

    if tasks:
        print(f"[*] Querying {len(tasks)} group(s) across "
              f"{len({d for d, _dc, _r, _s in scanned if _s in ('ok','unknown')})} "
              f"domain(s), {min(args.threads, len(tasks))} in parallel ...", file=sys.stderr)
        with ThreadPoolExecutor(max_workers=max(1, min(args.threads, len(tasks)))) as ex:
            for fut in as_completed([ex.submit(run_group, t) for t in tasks]):
                key, val = fut.result()
                findings[key] = val

    # Resolve every unique foreign SID in parallel.
    if not args.no_resolve:
        all_sids = set()
        for v in findings.values():
            all_sids.update(v.get("sids") or [])
        if all_sids:
            print(f"[*] Resolving {len(all_sids)} foreign principal(s) ...", file=sys.stderr)

            def run_resolve(sid):
                return sid, resolve_principal(sid, scanned, creds, args.nxc_bin, args.timeout)

            with ThreadPoolExecutor(max_workers=max(1, min(args.threads, len(all_sids)))) as ex:
                for fut in as_completed([ex.submit(run_resolve, s) for s in all_sids]):
                    sid, res = fut.result()
                    resolved[sid] = res

    import datetime
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    # per-domain counts for the summary
    counts = {}
    total = 0
    for dom, dc, realm, status in scanned:
        n = 0
        for grp in groups:
            f = findings.get((dom, grp), {})
            if not f.get("error"):
                n += len(f.get("sids") or [])
        counts[dom] = n
        total += n

    def label(kind):
        return {"group": "FOREIGN GROUP", "user": "FOREIGN USER"}.get(kind, "FOREIGN PRINCIPAL")

    def meaning(kind, grp, dom):
        if kind == "group":
            return (f"Every member of this external group effectively has "
                    f"'{grp}' rights in {dom}.")
        if kind == "user":
            return f"This external account has '{grp}' rights in {dom}."
        return f"This external principal is a member of '{grp}' in {dom}."

    sep = "=" * 78
    dash = "-" * 78
    L = []
    L += [sep, "FOREIGN SECURITY PRINCIPALS IN PRIVILEGED GROUPS", sep, ""]
    L += [
        "WHAT THIS REPORT SHOWS",
        "This report lists accounts and groups from OTHER domains or forests that are",
        "members of privileged groups in the domains checked below. Because these groups",
        "are highly privileged, anyone who controls one of these external members",
        "effectively holds that privileged access here -- so compromising one domain can",
        "lead to compromise of the others across the trust.",
        "",
        f"Generated:       {now}",
        f"Domains checked: {', '.join(d for d, _dc, _r, _s in scanned)}",
        f"Groups checked:  {', '.join(groups)}",
        "",
        dash, "SUMMARY", dash,
    ]
    for dom, dc, realm, status in scanned:
        if status not in ("ok", "unknown"):
            L.append(f"  {dom:<34} could not check ({status})")
        else:
            n = counts[dom]
            L.append(f"  {dom:<34} " +
                     (f"{n} foreign principal(s) found" if n else "none found"))
    L.append("")
    L.append(f"  TOTAL: {total} foreign principal(s) in privileged groups"
             if total else "  TOTAL: no foreign principals found in the checked groups")
    L += ["", dash, "FINDINGS", dash, ""]

    for dom, dc, realm, status in scanned:
        L.append(f"Domain: {dom}")
        if status not in ("ok", "unknown"):
            L += [f"  Could not check this domain ({status}).", ""]
            continue
        shown = False
        for grp in groups:
            f = findings.get((dom, grp), {})
            if f.get("error"):
                L.append(f"  Privileged group: {grp}")
                L.append(f"    Could not read this group ({f['error']}).")
                shown = True
                continue
            if not f.get("sids"):
                continue
            shown = True
            L.append(f"  Privileged group: {grp}")
            for s in f["sids"]:
                name, kind = resolved.get(s, (None, "unknown"))
                if name:
                    L.append(f"    - {label(kind):<17} {name}")
                    L.append(f"        {meaning(kind, grp, dom)}")
                else:
                    L.append(f"    - FOREIGN PRINCIPAL   (could not resolve name)")
                    L.append(f"        SID {s} -- its home domain is not in the list, so")
                    L.append(f"        the name/type could not be looked up. Add that domain to see it.")
        if not shown:
            L.append("  No foreign principals found in the checked groups.")
        L.append("")

    L += [dash, "RECOMMENDATION", dash,
          "  - Remove external users and groups from each privileged group.",
          "  - Use dedicated, least-privileged admin accounts within each domain",
          "    (no cross-domain nesting).",
          "  - Enable change-auditing and real-time alerts on privileged group membership.",
          "  - Review privileged group membership and trust paths on a schedule.",
          "",
          "Scope note: this lists DIRECT members of each group. Name/type resolution is",
          "limited to domains included in the list and reachable at scan time.",
          sep]
    text = "\n".join(L) + "\n"

    with open(args.out, "w") as fh:
        fh.write(text)
    sys.stdout.write(text)

    # ---- CSV (one row per finding, readable headers) ----
    if args.csv:
        def kind_word(k):
            return {"group": "Foreign Group", "user": "Foreign User"}.get(
                k, "Foreign Principal (type undetermined)")
        with open(args.csv, "w", newline="") as fcsv:
            w = csv.writer(fcsv)
            w.writerow(["Domain", "PrivilegedGroup", "PrincipalType",
                        "ForeignPrincipal", "WhatItMeans", "ForeignSID"])
            first = True
            for dom, dc, realm, status in scanned:
                rows = []
                for grp in groups:
                    f = findings.get((dom, grp), {})
                    if f.get("error") or not f.get("sids"):
                        continue
                    for s in f["sids"]:
                        name, kind = resolved.get(s, (None, "unknown"))
                        rows.append([dom, grp, kind_word(kind), name or "(unresolved)",
                                     meaning(kind, grp, dom), s])
                if rows:
                    if not first:
                        w.writerow([])
                    for r in rows:
                        w.writerow(r)
                    first = False
        print(f"[*] CSV written to {args.csv}", file=sys.stderr)

    if args.xlsx:
        import os
        try:
            if os.path.exists(args.xlsx):
                os.remove(args.xlsx)              # never leave a stale file behind
            write_xlsx(args.xlsx, scanned, groups, findings, resolved)
            print(f"[*] Excel report written to {args.xlsx}", file=sys.stderr)
        except ImportError:
            print("[!] Excel output needs openpyxl (pip install openpyxl); skipped.",
                  file=sys.stderr)
        except Exception as e:
            print(f"[!] Excel write FAILED ({type(e).__name__}: {e}). The .txt/.csv "
                  f"are correct; the .xlsx was NOT written (no stale file left).",
                  file=sys.stderr)
    print(f"[*] Report written to {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
